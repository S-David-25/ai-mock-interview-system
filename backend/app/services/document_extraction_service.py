import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

import docx
import pypdf

logger = logging.getLogger("document_extraction_service")


@dataclass
class ExtractedDocument:
    text: str
    source: str
    used_ocr: bool
    page_count: Optional[int]
    character_count: int
    confidence: Optional[float] = None
    warnings: List[str] = field(default_factory=list)


class DocumentExtractionService:
    """Native-first PDF/DOCX extraction with OCR fallback for scanned PDFs."""

    MIN_MEANINGFUL_CHARACTERS = 40
    MIN_MEANINGFUL_WORDS = 5

    @classmethod
    def extract_document(cls, file_path: Path, file_type: Optional[str] = None) -> ExtractedDocument:
        path = Path(file_path)
        extension = (file_type or path.suffix).lower().lstrip(".")
        if extension == "pdf":
            return cls._extract_pdf(path)
        if extension == "docx":
            return cls._extract_docx(path)
        raise ValueError(f"Unsupported document type: .{extension}")

    @classmethod
    def _extract_pdf(cls, path: Path) -> ExtractedDocument:
        native_error = None
        native_text = ""
        page_count = None
        try:
            reader = pypdf.PdfReader(str(path))
            page_count = len(reader.pages)
            native_text = cls._normalize_text(
                "\n".join(page.extract_text() or "" for page in reader.pages)
            )
        except Exception as exc:
            native_error = str(exc)

        if cls._is_usable(native_text):
            return ExtractedDocument(
                text=native_text,
                source="native_pdf",
                used_ocr=False,
                page_count=page_count,
                character_count=len(native_text),
            )

        warnings = []
        if native_error:
            warnings.append(f"Native PDF extraction failed: {native_error}")
        else:
            warnings.append("Native PDF extraction produced insufficient usable text.")
        ocr_document = cls._ocr_pdf(path, page_count=page_count, warnings=warnings)
        if not cls._is_usable(ocr_document.text):
            raise ValueError("PDF extraction produced insufficient usable text after OCR.")
        return ocr_document

    @classmethod
    def _extract_docx(cls, path: Path) -> ExtractedDocument:
        try:
            document = docx.Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
            text = cls._normalize_text("\n".join(parts))
        except Exception as exc:
            raise ValueError(f"Failed to extract text from DOCX: {exc}") from exc

        if not cls._is_usable(text):
            raise ValueError("DOCX extraction produced no usable text.")
        return ExtractedDocument(
            text=text,
            source="native_docx",
            used_ocr=False,
            page_count=None,
            character_count=len(text),
        )

    @classmethod
    def _ocr_pdf(cls, path: Path, page_count: Optional[int], warnings: List[str]) -> ExtractedDocument:
        try:
            import pypdfium2 as pdfium
            import pytesseract
        except ImportError as exc:
            raise ValueError(
                "PDF text extraction produced insufficient text and OCR is unavailable. "
                "Install pypdfium2, pytesseract, and the Tesseract OCR executable."
            ) from exc

        try:
            pdf = pdfium.PdfDocument(str(path))
            page_text: List[str] = []
            confidence_values: List[float] = []
            for index in range(len(pdf)):
                page = pdf[index]
                bitmap = page.render(scale=2.0)
                image = bitmap.to_pil()
                page_text.append(pytesseract.image_to_string(image))
                try:
                    data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
                    values = [float(value) for value in data.get("conf", []) if str(value).strip() not in {"", "-1"}]
                    confidence_values.extend(values)
                except Exception:
                    pass
                page.close()
            text = cls._normalize_text("\n\n".join(page_text))
            confidence = round(sum(confidence_values) / len(confidence_values), 1) if confidence_values else None
            return ExtractedDocument(
                text=text,
                source="ocr",
                used_ocr=True,
                page_count=page_count if page_count is not None else len(pdf),
                character_count=len(text),
                confidence=confidence,
                warnings=warnings,
            )
        except Exception as exc:
            raise ValueError(f"OCR fallback failed: {exc}") from exc

    @staticmethod
    def _normalize_text(text: str) -> str:
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        normalized = re.sub(r"[ \t]+", " ", normalized)
        normalized = re.sub(r"\n{3,}", "\n\n", normalized)
        return normalized.strip()

    @classmethod
    def _is_usable(cls, text: str) -> bool:
        if not text or len(text) < cls.MIN_MEANINGFUL_CHARACTERS:
            return False
        meaningful = len(re.findall(r"[A-Za-z0-9]", text))
        words = len(text.split())
        return meaningful >= cls.MIN_MEANINGFUL_CHARACTERS // 2 and words >= cls.MIN_MEANINGFUL_WORDS
