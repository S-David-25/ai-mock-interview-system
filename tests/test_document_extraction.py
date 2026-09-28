import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import docx
from fpdf import FPDF

from app.services.document_extraction_service import DocumentExtractionService, ExtractedDocument


def pdf_bytes(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    for line in text.splitlines():
        pdf.cell(w=190, h=8, txt=line, ln=1)
    output = pdf.output(dest="S")
    return bytes(output) if isinstance(output, (bytes, bytearray)) else output.encode("latin-1")


def docx_bytes(text: str) -> bytes:
    document = docx.Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class TestDocumentExtraction(unittest.TestCase):
    def test_text_pdf_uses_native_extraction_without_ocr(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "resume.pdf"
            path.write_bytes(pdf_bytes("Python FastAPI PostgreSQL React Docker\nExperience building REST API systems."))
            with patch.object(DocumentExtractionService, "_ocr_pdf", side_effect=AssertionError("OCR should not run")):
                result = DocumentExtractionService.extract_document(path)
        self.assertEqual(result.source, "native_pdf")
        self.assertFalse(result.used_ocr)
        self.assertIn("FastAPI", result.text)
        self.assertGreater(result.page_count, 0)

    def test_docx_uses_native_extraction(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "resume.docx"
            path.write_bytes(docx_bytes("Python, FastAPI, PostgreSQL, React, CNN, BERT\nBackend engineer experience."))
            result = DocumentExtractionService.extract_document(path)
        self.assertEqual(result.source, "native_docx")
        self.assertFalse(result.used_ocr)
        for term in ("Python", "FastAPI", "PostgreSQL", "CNN", "BERT"):
            self.assertIn(term, result.text)

    def test_unusable_pdf_invokes_ocr_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scanned.pdf"
            path.write_bytes(b"not-a-readable-text-pdf")
            ocr_result = ExtractedDocument(
                text="Python FastAPI PostgreSQL scanned resume content", source="ocr", used_ocr=True,
                page_count=1, character_count=48, confidence=91.0,
            )
            with patch.object(DocumentExtractionService, "_ocr_pdf", return_value=ocr_result) as ocr:
                result = DocumentExtractionService.extract_document(path)
        ocr.assert_called_once()
        self.assertTrue(result.used_ocr)
        self.assertEqual(result.source, "ocr")
        self.assertIn("PostgreSQL", result.text)

    def test_corrupt_docx_fails_with_controlled_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "broken.docx"
            path.write_bytes(b"broken")
            with self.assertRaisesRegex(ValueError, "Failed to extract text from DOCX"):
                DocumentExtractionService.extract_document(path)

    def test_unsupported_type_fails_with_controlled_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "resume.txt"
            path.write_text("Python FastAPI", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Unsupported document type"):
                DocumentExtractionService.extract_document(path)


if __name__ == "__main__":
    unittest.main()
