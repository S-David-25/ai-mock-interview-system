import unittest
import io
from pathlib import Path
import docx
from fpdf import FPDF
from starlette.testclient import TestClient
from app.main import app
from app.database.base import init_db, reset_db

def create_sample_pdf_bytes(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    for line in text.split('\n'):
        try:
            pdf.cell(w=200, h=10, txt=line, ln=1)
        except TypeError:
            pdf.cell(w=200, h=10, text=line)
    try:
        out = pdf.output(dest='S')
    except Exception:
        out = pdf.output()
    if isinstance(out, (bytes, bytearray)):
        return bytes(out)
    return out.encode('latin-1')

def create_sample_docx_bytes(text: str) -> bytes:
    doc = docx.Document()
    for p in text.split('\n'):
        if p.strip():
            doc.add_paragraph(p)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

class TestFileUploadEndpoints(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        reset_db()
        cls.client = TestClient(app)

        # Register User
        res = cls.client.post("/api/auth/register", json={
            "name": "Candidate Charlie",
            "email": "charlie@university.edu",
            "password": "Password123!",
            "confirm_password": "Password123!"
        })
        cls.token = res.json()["access_token"]
        cls.headers = {"Authorization": f"Bearer {cls.token}"}

    def test_01_general_interview_pdf_resume_upload(self):
        res_create = self.client.post("/api/interviews", json={"interview_type": "general"}, headers=self.headers)
        int_id = res_create.json()["id"]

        pdf_bytes = create_sample_pdf_bytes(
            "Resume: Charlie Candidate\nEducation: BS Computer Science\nSkills: Python, FastAPI, React, SQL, Algorithms\nExperience: 2 years software engineering internship."
        )

        files = {"file": ("charlie_resume.pdf", pdf_bytes, "application/pdf")}
        res_upload = self.client.post(f"/api/interviews/{int_id}/upload-resume", files=files, headers=self.headers)
        self.assertEqual(res_upload.status_code, 200)
        data = res_upload.json()
        self.assertEqual(data["file_type"], "resume")
        self.assertEqual(data["original_name"], "charlie_resume.pdf")
        self.assertGreater(data["word_count"], 10)
        self.assertEqual(data["interview_status"], "ready")

        # Verify status endpoint
        res_status = self.client.get(f"/api/interviews/{int_id}/status", headers=self.headers)
        self.assertEqual(res_status.status_code, 200)
        self.assertTrue(res_status.json()["is_ready"])
        self.assertTrue(res_status.json()["is_resume_uploaded"])

    def test_02_company_interview_full_flow(self):
        res_create = self.client.post("/api/interviews", json={
            "interview_type": "company",
            "company_name": "Amazon",
            "job_role": "Backend Engineer"
        }, headers=self.headers)
        int_id = res_create.json()["id"]

        # Step 1: Upload Resume (DOCX)
        docx_bytes = create_sample_docx_bytes(
            "Resume: Charlie Candidate\nSenior Full Stack Developer\nExpertise in distributed systems, microservices, AWS, Docker."
        )
        files_resume = {"file": ("charlie_cv.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res_res = self.client.post(f"/api/interviews/{int_id}/upload-resume", files=files_resume, headers=self.headers)
        self.assertEqual(res_res.status_code, 200)

        # Status should still be 'setup' because JD is missing for company interview
        res_status_mid = self.client.get(f"/api/interviews/{int_id}/status", headers=self.headers)
        self.assertFalse(res_status_mid.json()["is_ready"])
        self.assertIn("job description", res_status_mid.json()["message"].lower())

        # Step 2: Upload JD (PDF)
        jd_pdf_bytes = create_sample_pdf_bytes(
            "Job Description: Backend Engineer at Amazon\nRequirements: 3+ years experience with REST APIs, high scalability, databases."
        )
        files_jd = {"file": ("amazon_jd.pdf", jd_pdf_bytes, "application/pdf")}
        res_jd = self.client.post(f"/api/interviews/{int_id}/upload-jd", files=files_jd, headers=self.headers)
        self.assertEqual(res_jd.status_code, 200)
        self.assertEqual(res_jd.json()["interview_status"], "ready")

        # Now status should be ready
        res_status_final = self.client.get(f"/api/interviews/{int_id}/status", headers=self.headers)
        self.assertTrue(res_status_final.json()["is_ready"])
        self.assertTrue(res_status_final.json()["is_resume_uploaded"])
        self.assertTrue(res_status_final.json()["is_jd_uploaded"])

    def test_03_unsupported_file_extension(self):
        res_create = self.client.post("/api/interviews", json={"interview_type": "general"}, headers=self.headers)
        int_id = res_create.json()["id"]

        fake_file = io.BytesIO(b"malicious script or invalid format")
        files = {"file": ("malicious.exe", fake_file, "application/octet-stream")}
        res = self.client.post(f"/api/interviews/{int_id}/upload-resume", files=files, headers=self.headers)
        self.assertEqual(res.status_code, 400)
        self.assertIn("unsupported file format", res.json()["detail"].lower())

    def test_04_reject_jd_on_general_interview(self):
        res_create = self.client.post("/api/interviews", json={"interview_type": "general"}, headers=self.headers)
        int_id = res_create.json()["id"]

        pdf_bytes = create_sample_pdf_bytes("Sample Job Description")
        files = {"file": ("jd.pdf", pdf_bytes, "application/pdf")}
        res = self.client.post(f"/api/interviews/{int_id}/upload-jd", files=files, headers=self.headers)
        self.assertEqual(res.status_code, 400)
        self.assertIn("company-specific", res.json()["detail"].lower())

    def test_05_company_interview_pasted_jd_only(self):
        """Case 1: Paste JD text only without any file upload."""
        res_create = self.client.post("/api/interviews", json={
            "interview_type": "company",
            "company_name": "Google",
            "job_role": "Backend Engineer"
        }, headers=self.headers)
        int_id = res_create.json()["id"]

        # Upload resume
        resume_pdf = create_sample_pdf_bytes(
            "Resume: Charlie Candidate\nEducation: BS Computer Science\nSkills: Python, FastAPI, Docker, SQL"
        )
        self.client.post(f"/api/interviews/{int_id}/upload-resume", files={"file": ("resume.pdf", resume_pdf, "application/pdf")}, headers=self.headers)

        # Submit pasted JD text
        jd_text = (
            "Job Title: Backend Engineer at Google\n"
            "Requirements: 3+ years experience with Python, FastAPI, PostgreSQL, Kubernetes, Docker.\n"
            "Strong team collaboration and technical problem solving."
        )
        res_submit = self.client.post(f"/api/interviews/{int_id}/submit-jd", data={"jd_text": jd_text}, headers=self.headers)
        self.assertEqual(res_submit.status_code, 200)
        self.assertEqual(res_submit.json()["interview_status"], "ready")

        # Verify status endpoint
        res_status = self.client.get(f"/api/interviews/{int_id}/status", headers=self.headers)
        self.assertEqual(res_status.status_code, 200)
        self.assertTrue(res_status.json()["is_ready"])
        self.assertTrue(res_status.json()["is_jd_uploaded"])
        self.assertTrue(res_status.json()["is_resume_uploaded"])

    def test_06_company_interview_jd_docx_upload(self):
        """Case 2: Upload JD DOCX file only."""
        res_create = self.client.post("/api/interviews", json={
            "interview_type": "company",
            "company_name": "Microsoft",
            "job_role": "Cloud Architect"
        }, headers=self.headers)
        int_id = res_create.json()["id"]

        # Upload resume
        resume_pdf = create_sample_pdf_bytes("Resume: Charlie Candidate\nSkills: Azure, Python, Kubernetes")
        self.client.post(f"/api/interviews/{int_id}/upload-resume", files={"file": ("resume.pdf", resume_pdf, "application/pdf")}, headers=self.headers)

        # Upload JD docx
        jd_docx = create_sample_docx_bytes(
            "Job Description: Cloud Architect at Microsoft\nRequirements: 5+ years experience with Azure, distributed systems, CI/CD."
        )
        res_jd = self.client.post(
            f"/api/interviews/{int_id}/upload-jd",
            files={"file": ("ms_jd.docx", jd_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=self.headers
        )
        self.assertEqual(res_jd.status_code, 200)
        self.assertEqual(res_jd.json()["interview_status"], "ready")

        res_status = self.client.get(f"/api/interviews/{int_id}/status", headers=self.headers)
        self.assertTrue(res_status.json()["is_ready"])
        self.assertTrue(res_status.json()["is_jd_uploaded"])

    def test_07_company_interview_both_pasted_and_pdf(self):
        """Case 3: Paste JD text AND upload PDF file."""
        res_create = self.client.post("/api/interviews", json={
            "interview_type": "company",
            "company_name": "Netflix",
            "job_role": "Distributed Systems Engineer"
        }, headers=self.headers)
        int_id = res_create.json()["id"]

        # Upload resume
        resume_pdf = create_sample_pdf_bytes("Resume: Charlie Candidate\nSkills: Python, Java, Kafka, Microservices")
        self.client.post(f"/api/interviews/{int_id}/upload-resume", files={"file": ("resume.pdf", resume_pdf, "application/pdf")}, headers=self.headers)

        # 1. Submit pasted JD text
        pasted_text = "Role Overview: Distributed Systems Engineer for content delivery.\nMust have strong concurrency and distributed algorithms."
        res_submit = self.client.post(f"/api/interviews/{int_id}/submit-jd", data={"jd_text": pasted_text}, headers=self.headers)
        self.assertEqual(res_submit.status_code, 200)

        # 2. Upload JD PDF file
        file_text = "Job Description: Netflix Engineering\nRequirements: Experience with Java, gRPC, Cassandra, high throughput systems."
        pdf_bytes = create_sample_pdf_bytes(file_text)
        res_upload = self.client.post(f"/api/interviews/{int_id}/upload-jd", files={"file": ("netflix_jd.pdf", pdf_bytes, "application/pdf")}, headers=self.headers)
        self.assertEqual(res_upload.status_code, 200)

        # Verify status endpoint
        res_status = self.client.get(f"/api/interviews/{int_id}/status", headers=self.headers)
        self.assertTrue(res_status.json()["is_ready"])
        self.assertTrue(res_status.json()["is_jd_uploaded"])

    def test_08_company_interview_both_pasted_and_docx(self):
        """Case 3: Paste JD text AND upload DOCX file."""
        res_create = self.client.post("/api/interviews", json={
            "interview_type": "company",
            "company_name": "Apple",
            "job_role": "iOS Developer"
        }, headers=self.headers)
        int_id = res_create.json()["id"]

        # Upload resume
        resume_pdf = create_sample_pdf_bytes("Resume: Charlie Candidate\nSkills: Swift, SwiftUI, Objective-C, iOS SDK")
        self.client.post(f"/api/interviews/{int_id}/upload-resume", files={"file": ("resume.pdf", resume_pdf, "application/pdf")}, headers=self.headers)

        # 1. Submit pasted JD text
        pasted_text = "Role: iOS Developer\nRequirements: Swift, SwiftUI, Combine framework, Core Data."
        res_submit = self.client.post(f"/api/interviews/{int_id}/submit-jd", data={"jd_text": pasted_text}, headers=self.headers)
        self.assertEqual(res_submit.status_code, 200)

        # 2. Upload JD DOCX file
        docx_bytes = create_sample_docx_bytes("Apple Job Description: Mobile Engineering team. App Store deployment experience.")
        res_upload = self.client.post(
            f"/api/interviews/{int_id}/upload-jd",
            files={"file": ("apple_jd.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=self.headers
        )
        self.assertEqual(res_upload.status_code, 200)

        res_status = self.client.get(f"/api/interviews/{int_id}/status", headers=self.headers)
        self.assertTrue(res_status.json()["is_ready"])
        self.assertTrue(res_status.json()["is_jd_uploaded"])

    def test_09_company_interview_empty_jd_text_rejected(self):
        """Reject empty or whitespace-only pasted JD text."""
        res_create = self.client.post("/api/interviews", json={
            "interview_type": "company",
            "company_name": "Oracle",
            "job_role": "Database Engineer"
        }, headers=self.headers)
        int_id = res_create.json()["id"]

        res_empty = self.client.post(f"/api/interviews/{int_id}/submit-jd", data={"jd_text": "   "}, headers=self.headers)
        self.assertEqual(res_empty.status_code, 400)
        self.assertIn("cannot be empty", res_empty.json()["detail"].lower())

if __name__ == "__main__":
    unittest.main()
